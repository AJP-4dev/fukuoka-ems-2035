import sys
from pptx import Presentation
from pptx.util import Emu
prs=Presentation(sys.argv[1]); SW,SH=prs.slide_width,prs.slide_height
IN=914400; bad=0
for i,sl in enumerate(prs.slides,1):
    print(f"--- slide {i}")
    boxes=[]
    for sh in sl.shapes:
        x,y,w,h=[v/IN for v in (sh.left,sh.top,sh.width,sh.height)]
        kind=sh.shape_type
        t=sh.text_frame.text[:18].replace("\n","/") if sh.has_text_frame else ("chart" if sh.has_chart else "tbl" if sh.has_table else "")
        sizes=[r.font.size.pt for p in sh.text_frame.paragraphs for r in p.runs if r.font.size] if sh.has_text_frame else []
        print(f"  {x:5.2f} {y:5.2f} {w:5.2f} {h:5.2f} {str(kind)[:14]:14} {t} {min(sizes) if sizes else ''}")
        if x<-0.001 or y<-0.001 or x+w>SW/IN+0.001 or y+h>SH/IN+0.001: print("   !! OUT OF BOUNDS"); bad+=1
        if sh.has_text_frame and sh.text_frame.text.strip() and not sh.has_chart:
            boxes.append((sh,x,y,w,h))
        if sizes and min(sizes)<10: print("   !! font <10"); bad+=1
    for a in range(len(boxes)):
        for b in range(a+1,len(boxes)):
            _,x1,y1,w1,h1=boxes[a];_,x2,y2,w2,h2=boxes[b]
            ix=max(0,min(x1+w1,x2+w2)-max(x1,x2));iy=max(0,min(y1+h1,y2+h2)-max(y1,y2))
            ar=ix*iy; m=min(w1*h1,w2*h2)
            if m>0 and ar/m>0.10: print(f"   !! overlap {ar/m:.0%}: '{boxes[a][0].text_frame.text[:14]}' vs '{boxes[b][0].text_frame.text[:14]}'"); bad+=1
    # chart vs text overlap
    for sh in sl.shapes:
        if sh.has_chart:
            cx,cy,cw,ch=[v/IN for v in (sh.left,sh.top,sh.width,sh.height)]
            for (o,x,y,w,h) in boxes:
                ix=max(0,min(cx+cw,x+w)-max(cx,x));iy=max(0,min(cy+ch,y+h)-max(cy,y))
                if ix*iy>0.05: print(f"   !! chart overlaps text '{o.text_frame.text[:14]}'"); bad+=1
print("issues:",bad)

# heuristic text-fit: full-width CJK ~ 1.0em, ASCII ~0.55em; line height 1.2em
import math
print("== text fit heuristic")
for i,sl in enumerate(prs.slides,1):
    for sh in sl.shapes:
        if not sh.has_text_frame or not sh.text_frame.text.strip(): continue
        w=sh.width/IN; h=sh.height/IN
        need=0
        for p in sh.text_frame.paragraphs:
            sz=max([r.font.size.pt for r in p.runs if r.font.size] or [16])
            em=sz/72
            L=sum(1.0*em if ord(c)>0x2E80 else 0.55*em for c in "".join(r.text for r in p.runs))
            lines=max(1,math.ceil(L/max(w-0.02,0.1)))
            need+=lines*em*1.2
        if need>h+0.05: print(f"  slide {i}: need {need:.2f} > h {h:.2f} w {w:.2f}: {sh.text_frame.text[:24]!r}")
