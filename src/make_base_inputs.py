"""基礎入力 base_inputs.json を src/raw/ のExcel/CSVから生成する。

  .venv/bin/python -I src/make_base_inputs.py
"""
import csv
import json
import pathlib
import re
import unicodedata

import openpyxl
import xlrd

HERE = pathlib.Path(__file__).resolve().parent
RAW = HERE / 'raw'
OUT = HERE / 'base_inputs.json'

# 区の並び（登録人口・推計人口Excelの列順）と区コード（国土数値情報 N03 / 総務省）
WARDS = ['東', '博多', '中央', '南', '城南', '早良', '西']
CODES = {'東': '40131', '博多': '40132', '中央': '40133', '南': '40134',
         '西': '40135', '城南': '40136', '早良': '40137'}
BANDS = [f'{a}～{a + 4}歳' for a in range(0, 90, 5)] + ['90歳以上']

BASE = 'https://www.city.fukuoka.lg.jp/soki/tokeichosa/shisei/toukei/'
JUKI = BASE + 'jinkou/tourokujinkou/'


def nfkc(s):
    t = unicodedata.normalize('NFKC', str(s)).replace(' ', '').replace('　', '')
    return t.replace('~', '～')   # NFKCは全角チルダを半角にするので戻す


def sheet_rows(path, sheet_match):
    """(ヘッダ付き) シートを行リストで返す。xlsx/xls 両対応。"""
    path = pathlib.Path(path)
    if path.suffix == '.xlsx':
        wb = openpyxl.load_workbook(path, data_only=True)
        name = next(n for n in wb.sheetnames if sheet_match(nfkc(n)))
        return [list(r) for r in wb[name].iter_rows(values_only=True)], name
    wb = xlrd.open_workbook(path)
    name = next(n for n in wb.sheet_names() if sheet_match(nfkc(n)))
    s = wb.sheet_by_name(name)
    return [s.row_values(i) for i in range(s.nrows)], name


def band_index(label):
    """行ラベルから 5歳階級の下限年齢を返す（5歳階級行でなければ None）。"""
    t = nfkc(label)
    m = re.fullmatch(r'(\d+)～(\d+)歳?', t)
    if m and int(m.group(2)) - int(m.group(1)) == 4:
        return int(m.group(1))
    m = re.fullmatch(r'(\d+)歳以上', t)
    if m and int(m.group(1)) >= 100:
        return int(m.group(1))
    return None


def parse_age_table(rows, units):
    """rows から {unit: {19帯: 人数}} を作る。units={'福岡市':'市', '東区':'東',...}。
    総数行との突き合わせ（19帯合計==総数）を検証する。"""
    hdr = next(r for r in rows if r and nfkc(r[0]) == '年齢' and any(nfkc(c or '') == '福岡市' for c in r))
    col = {}
    for i, c in enumerate(hdr):
        if c and nfkc(c) in units and units[nfkc(c)] not in col:
            col[units[nfkc(c)]] = i
    assert set(col) == set(units.values()), (col, units)
    total = next(r for r in rows if r and nfkc(r[0]) == '総数')
    raw = {k: {} for k in col}
    seen_hdr = False
    for r in rows:
        if r is hdr:
            seen_hdr = True
        if not seen_hdr or not r or r[0] is None:
            continue
        b = band_index(r[0])
        if b is None or b in raw['市']:
            continue            # 最初の出現のみ（各歳行・重複見出しを除く）
        for k, ci in col.items():
            raw[k][b] = int(r[ci])
    out = {}
    for k, d in raw.items():
        assert all(b in d for b in range(0, 100, 5)), (k, sorted(d))
        vals = {BANDS[i]: d[a] for i, a in enumerate(range(0, 90, 5))}
        vals['90歳以上'] = sum(v for a, v in d.items() if a >= 90)
        assert sum(vals.values()) == int(total[col[k]]), (k, sum(vals.values()), total[col[k]])
        out[k] = vals
    return out


UNITS = {'福岡市': '市', **{f'{w}区': w for w in WARDS}}


def city_series():
    res = {}
    for y in range(2013, 2025):
        f = RAW / (f'{y}nennreibetu-sousuu.xls' if y == 2013 else f'{y}nenreibetu-sousuu.xlsx')
        rows, name = sheet_rows(f, lambda n: '9月末' in n)
        res[str(y)] = parse_age_table(rows, UNITS)['市']
    return res


def pop2025():
    """登録人口 住民基本台帳（日本人＋外国人）2024-12-31（2025-01-01 に最も近い公表値）"""
    rows, name = sheet_rows(RAW / '2024nenreibetu-sousuu.xlsx', lambda n: '12月末' in n)
    return parse_age_table(rows, UNITS)


def commute():
    s = xlrd.open_workbook(RAW / 'toukeisyo2024_0319.xls').sheet_by_index(0)
    out = {}
    for i in range(s.nrows):
        n = nfkc(s.cell_value(i, 0))
        if n in {f'{w}区' for w in WARDS}:
            r = s.row_values(i)
            night, day = int(r[2]), int(r[11])
            inflow = int(r[12] + r[13] + r[14])         # 他区・他市町村・他県に常住する従業者・通学者
            outflow = int(r[7] + r[8] + r[9])           # 他区・他市町村・他県で従業・通学
            assert day == night - outflow + inflow, (n, day, night, outflow, inflow)
            out[n[:-1]] = {'night20': night, 'day20': day, 'in': inflow, 'out': outflow,
                           'dn': round(day / night, 3)}
    assert set(out) == set(WARDS)
    return out


def areas():
    res = {}
    with open(RAW / 'R8_07_mencho.csv', encoding='utf-8-sig') as f:
        for r in csv.reader(f):
            if len(r) > 4 and r[0] in CODES.values():
                res[r[0]] = round(float(r[4]), 2)
    return res


def main():
    p25, cm, ar = pop2025(), commute(), areas()
    wards = {}
    for w in WARDS:
        wards[w] = {'code': CODES[w], 'pop2025': p25[w],
                    'area': ar[CODES[w]], **cm[w]}
    out = {
        'wards': wards,
        'cityAgeByYear': city_series(),
        'sources': {
            'asof': {
                'pop2025': '2024-12-31',
                'pop2025_note': '登録人口(住基,日本人+外国人)。2025-01-01時点の表は無く、直近の2024年12月末(R6.12月末)を採用。2025-01-31(1月末)も存在するが12月末のほうが1日差で近い',
                'cityAgeByYear': '各年9月30日（登録人口 日本人+外国人 の9月末シート）',
                'day_night': '2020-10-01 国勢調査',
                'area': '令和6年(2024)全国都道府県市区町村別面積調',
            },
            'pop2025': JUKI + 'TourokuJinko_kubetsu.html  -> documents/2024nenreibetu-sousuu.xlsx (シート R6.12月末)',
            'cityAgeByYear': JUKI + 'TourokuJinko_kubetsu.html  -> documents/{2014..2024}nenreibetu-sousuu.xlsx, 2013nennreibetu-sousuu.xls (シート 9月末)',
            'day_night': BASE + 'toukeisyo/2024/toukeisyo2024-3.html -> documents/0319.xls (第19表, 令和2年国勢調査 従業地・通学地集計)',
            'day_night_def': 'night20=常住人口, day20=昼間人口, in=他区・他市町村・他県に常住する従業・通学者, out=他区・他市町村・他県への従業・通学者(従業地不詳を除く)。day20=night20-out+in が成立。dn=day20/night20(3桁)',
            'area': 'https://www.gsi.go.jp/KOKUJYOHO/MENCHO-title.htm -> https://www.gsi.go.jp/KOKUJYOHO/MENCHO/backnumber/R8_07_mencho.csv (令和6年1月1日列; 福岡市各区の値は令和6年中不変)',
            'codes': '東40131 博多40132 中央40133 南40134 西40135 城南40136 早良40137',
        },
    }
    OUT.write_text(json.dumps(out, ensure_ascii=False, indent=1), encoding='utf-8')
    tot = sum(sum(wards[w]['pop2025'].values()) for w in WARDS)
    print('pop2025 total', tot)


if __name__ == '__main__':
    main()
